"""
The QA tracker, answered row by row.

Written as a script rather than a hand-made spreadsheet so the sheet can be
rebuilt when the next round of rows arrives, and so the evidence column says
where each answer was actually checked rather than being asserted once and
forgotten.

Status vocabulary, used strictly:
  Done          — built, tested, deployed, and seen working on the live site
  Done (code)   — built and tested, but the live check needs data this book
                  does not have (a configured custom field, say)
  Partly        — some of the row is done and the rest needs something from you
  Needs detail  — cannot be actioned without the screenshot or link the row
                  refers to
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROWS = [
    # sl, module, page, issue, recommendation (theirs), their status, our status, what was done, evidence
    (
        18, "Sales", "Invoices",
        "Invoice creation page",
        "Custom fields should appear below the Customer/Reference section.",
        "Pending", "Done (code)",
        "Custom fields render directly under the whole header band — customer, number, dates and refs — "
        "ruled off from it, two per row. They were already there before this round.",
        "src/features/sales/index.jsx — the block guarded by hasCustomFieldsAt(customFields,'header','reference'). "
        "Not visible on the live book because this company has defined no custom fields.",
    ),
    (
        27, "Sales", "Invoices",
        "Export invoices shows pdf/excel/csv directly under More; it should show only when Export is clicked.",
        "Options under More: Import invoices, Export invoices, Invoice settings, Custom fields.",
        "Pending", "Done",
        "The three formats sit behind one Export entry and open only when it is pressed. The same collapsed "
        "picker now applies to all 24 lists, not just invoices — a test fails on the next list that ships a bare Export.",
        "Live: Invoices › More shows Import invoices · Export invoices (collapsed) · Invoice settings · Custom fields · "
        "Invoice template · Recurring. Test: src/test/export-asks-the-format.test.js",
    ),
    (
        29, "Sales", "Invoices",
        "Headings in sales invoice page override the filter and search option (image link attached).",
        "See attached image.",
        "Pending", "Needs detail",
        "Could not reproduce. The column headings stick inside their own scroll container and the page header is "
        "not sticky, so there is no overlap in any width tried. The row refers to an image link that did not come through.",
        "Send the screenshot and the window width; I will fix it the same day.",
    ),
    (
        32, "Sales", "Receipts",
        "UI & Wiring",
        "Change the UI as per the link given, in addition to the existing form. TDS, bank charges, other charges "
        "entries implemented in receipts.",
        "Pending", "Partly",
        "TDS, bank charges and other charges are implemented — a Deductions band on the receipt, each posting to its "
        "own ledger line, with the summary showing what actually reaches the account. The 'UI as per the link' half "
        "cannot be done: the link is not in the sheet.",
        "Live: Receipts › Record Receipt. Send the reference link for the layout half.",
    ),
    (
        34, "Sales", "Estimates / Quotes",
        "Creation or view page should be the same as invoice creation and invoice page. Numbering and requirements "
        "of quotation should be different.",
        "Same layout as invoice; own numbering.",
        "Pending", "Done",
        "Quotation carries the invoice's two ruled header columns, the same nine-column line grid (unit, discount, "
        "tax rate), the same add control and hint, totals block, amount in words and pinned running total. Numbering "
        "stays its own series, and the second date is Valid Until because a quotation expires rather than falls due.",
        "Live: Sales › Quotations › New Quotation. Tests: src/features/sales/docFormParity.test.jsx",
    ),
    (
        35, "Sales", "Sales orders",
        "Creation or view page should be the same as invoice creation and invoice page. Numbering and requirements "
        "of sales orders should be different.",
        "Same layout as invoice; own numbering.",
        "Pending", "Done",
        "Same treatment as the quotation. The form also stopped opening inside the list — it replaces it, the way an "
        "invoice does — and the first total is called Subtotal rather than Taxable value.",
        "Live: Sales › Sales Orders › New Sales Order. Tests: docFormParity.test.jsx ('a document form is a screen').",
    ),
    (
        36, "Sales", "Credit note",
        "UI & Wiring — see attached screenshot.",
        "See attached screenshot.",
        "Pending", "Done",
        "Sales Returns rebuilt to the invoice layout. Behind the layout was a defect: the credit-note grid had no "
        "discount column while the invoice it reverses does, so a note against a discounted line returned more than "
        "had been charged. That is fixed.",
        "Live: Sales › Sales Returns › New Credit Note. If your screenshot shows something else, send it.",
    ),
    (
        37, "Sales", "Recurring invoices",
        "UI & Wiring — see attached screenshot.",
        "See attached screenshot.",
        "Pending", "Done",
        "Form: actions pinned at the head like every other document, description and unit added to the grid, a running "
        "'each run' figure, and the customer picker no longer splits its label and field across two grid columns. "
        "List: shared filter band and the same per-column filter and sort as the invoice list.",
        "Live: Sales › Recurring › New Schedule, and the list behind it.",
    ),
    (
        48, "Sales", "Invoices",
        "Numbers appear in bold everywhere; should be simple and normal.",
        "Change the UI, keep it simple and normal.",
        "New", "Done",
        "Headline figures, stat-card values and document totals were all half-bold, so a page of figures read as a page "
        "of emphasis. All normal weight now — emphasis comes from size and colour. Pinned by a test.",
        "Live: any list or form. Test: src/test/quiet-emphasis.test.js",
    ),
    (
        49, "Sales", "Invoices",
        "Record receipt from the invoice row should open the receipt screen prefilled with that invoice's data "
        "(amount, customer, invoice) and still be editable.",
        "Prefill from the invoice; allow changes.",
        "New", "Done",
        "Amount outstanding, the invoice reference and the allocation were already prefilled. The customer was not — "
        "an invoice that came back from the server carries the customer's name and no local id. The name is now matched "
        "to the master, so all four prefill and every field stays editable.",
        "Live: Invoices › row menu › Record Receipt. Test: src/features/payments/recordReceiptRouting.test.jsx",
    ),
    (
        50, "Sales", "Invoices",
        "Navigating to another page from the current page should offer a back button to return.",
        "Add a back button on the destination.",
        "New", "Done",
        "The bar existed but always said 'Back to Home' whatever screen sent you there. It now names the screen it "
        "returns to, taken from the rail so the two cannot disagree.",
        "Live: Invoices › Record Receipt shows '← Back to Invoices'.",
    ),
    (
        51, "Sales", "Invoices",
        "Move the TDS selection below the invoice total. After selection show TDS only, not the section; while "
        "selecting show section, rate, everything.",
        "Fold the chooser after selection.",
        "New", "Done",
        "TDS sits under the total. Closed it is a single '+ TDS deduction' line; opened it shows section, rate and "
        "payee; once chosen it folds away and the totals read 'Less: TDS' with a Change beside it. A first pass left "
        "the opener and the chooser on screen together — corrected and pinned by two tests.",
        "Live: Invoices › New Invoice, below Total. Tests: docFormParity.test.jsx (TDS is chosen once).",
    ),
    (
        52, "Sales", "Receipts",
        "Record receipt is completely off — redo the whole thing.",
        "Redo.",
        "New", "Partly",
        "Rebuilt: two ruled header columns, deductions band, receipt summary, outstanding-invoice allocation, "
        "mandatory-field marks, one action bar at the head and the net figure pinned at the foot. Whether that is what "
        "'completely off' meant, I cannot tell from the row.",
        "Live: Sales › Receipts › Record Receipt. Tell me what is still wrong and I will take it point by point.",
    ),
]

STATUS_FILL = {
    "Done": ("1B7F4B", "E8F5EE"),
    "Done (code)": ("1B6F7F", "E6F3F6"),
    "Partly": ("9A6A00", "FDF3DF"),
    "Needs detail": ("A33A3A", "FBEAEA"),
}

HEADERS = [
    ("Sl no", 7),
    ("Module", 10),
    ("Page", 16),
    ("Issue (as raised)", 42),
    ("Recommendation (as raised)", 42),
    ("Your status", 12),
    ("Our status", 14),
    ("What was done", 62),
    ("Where to verify", 52),
]

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def build(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "QA tracker — status"

    ws["A1"] = "Neev One — QA tracker, answered"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        "Every row below was checked on the live site at https://68.233.107.83.sslip.io unless the evidence column "
        "says otherwise. 603 tests pass, 0 lint errors."
    )
    ws["A2"].font = Font(size=10, color="595959")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS))

    head_row = 4
    for col, (title, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=head_row, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[head_row].height = 26

    for i, row in enumerate(ROWS):
        r = head_row + 1 + i
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
        status = row[6]
        ink, bg = STATUS_FILL[status]
        sc = ws.cell(row=r, column=7)
        sc.font = Font(bold=True, color=ink)
        sc.fill = PatternFill("solid", fgColor=bg)
        ws.cell(row=r, column=6).alignment = Alignment(vertical="top", horizontal="center")
        ws.row_dimensions[r].height = 74

    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(HEADERS))}{head_row + len(ROWS)}"

    # A second sheet: what else changed in the same rounds, so the sheet is a
    # fair account rather than only the rows that were raised.
    extra = wb.create_sheet("Also fixed (not raised)")
    extra_headers = ["Area", "What was wrong", "Why it mattered"]
    for col, title in enumerate(extra_headers, start=1):
        c = extra.cell(row=1, column=col, value=title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = border
    for w, col in zip((26, 62, 62), "ABC"):
        extra.column_dimensions[col].width = w

    extras = [
        ("Import", "Import failed at the last step with 'Data import is switched off for this company'.",
         "The feature is opt-in. You chose a type, pasted a file and pressed Check before anything said so. The screen "
         "now says it up front and names where to switch it on."),
        ("Import templates", "The template downloaded immediately with every column on it.",
         "Downloading now asks which columns, with select-all and required-only. Required columns cannot be unticked."),
        ("Customer / Vendor", "The address and contact tables pushed Actions off the right edge; the delete icon was unreachable.",
         "Fixed-pixel inputs forced the table past the dialog. Inputs are fluid, Actions is pinned to the right."),
        ("Vendor master", "A vendor entered on the Vendors screen never reached the server.",
         "It lived in one browser; signing in elsewhere came back missing those vendors, with bills pointing at no party."),
        ("Vendor master", "Credit Limit was on screen and bound to nothing.",
         "Whatever was typed was dropped on save."),
        ("Vendor master", "A new vendor's ledger opened Dr.",
         "A vendor balance is money owed — it opens Cr. Dr is the customer's default."),
        ("Bills", "The bill's line grid had no discount column.",
         "A supplier discount had to be back-worked into the rate, which is how a bill stops matching its order."),
        ("Sales returns", "The credit-note grid had no discount column.",
         "A note against a discounted line returned more than had been charged."),
        ("Deploy", "A deploy deleted the old bundle as it wrote the new index.html.",
         "A page loaded mid-deploy asked for a file that had just gone — a blank screen and a 404. Three passes now."),
        ("Settings", "Settings expanded to 34 entries inside the global sidebar.",
         "Redesigned as its own workspace: a hub of six categories, local navigation, breadcrumbs and search."),
    ]
    for i, row in enumerate(extras, start=2):
        for col, value in enumerate(row, start=1):
            c = extra.cell(row=i, column=col, value=value)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = border
        extra.row_dimensions[i].height = 46
    extra.freeze_panes = "A2"

    wb.save(path)
    return path


if __name__ == "__main__":
    out = build("docs/NEEV-QA-TRACKER-STATUS.xlsx")
    print(f"wrote {out}")
